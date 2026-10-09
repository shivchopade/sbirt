from __future__ import annotations

from collections import Counter
from pathlib import Path

from app.core.config import Settings
from app.services.extraction import extract
from app.services.extraction.contact import extract_contact
from app.services.preprocessing import preprocess_document
from app.services.preprocessing.normalize import canonical, normalize_text
from app.services.visibility.engine import apply_visibility


def test_lexical_text_stopwords_and_lemmas(nlp_engine) -> None:
    a = nlp_engine.analyze(["Built REST APIs with Go and deployed the services using Docker containers"])[0]
    assert "build" in a.lexical and "deploy" in a.lexical
    assert {"skill:go", "skill:rest_api", "skill:docker"} <= set(a.lexical)
    assert not {"the", "with", "and", "using"} & set(a.lexical)
    assert a.has_verb


def test_skill_aliases(nlp_engine) -> None:
    a = nlp_engine.analyze(["Strong in js, postgres and ml; ran k8s clusters on amazon web services"])[0]
    assert {"javascript", "postgresql", "machine learning", "kubernetes", "aws"} <= set(a.skills)


def test_ambiguous_skills_need_exact_spelling(nlp_engine) -> None:
    yes = nlp_engine.analyze(["Wrote services in Go and dashboards in React"])[0]
    no = nlp_engine.analyze(["Ready to go the extra mile and react quickly to incidents"])[0]
    assert {"go", "react"} <= set(yes.skills)
    assert not ({"go", "react"} & set(no.skills))


def test_skill_phrases_survive_stopword_removal(nlp_engine) -> None:
    """'Go' and 'IT'-like tokens are spaCy stopwords; matched skills must still be kept."""
    a = nlp_engine.analyze(["Experience with Go and C++ and C#"])[0]
    assert {"skill:go", "skill:c++", "skill:c#"} <= set(a.lexical)


def test_repetition_adds_no_skills(nlp_engine) -> None:
    once = nlp_engine.analyze(["Python developer"])[0]
    many = nlp_engine.analyze(["Python " * 30])[0]
    assert set(once.skills) == set(many.skills) == {"python"}


def test_pii_masking_keeps_skills(nlp_engine) -> None:
    text = "Aarav Sharma led Django and Kafka migrations; mail aarav.sharma@example.com, call +91 98765 43210, see github.com/aarav"
    a = nlp_engine.analyze([text], mask_names=("Aarav Sharma",))[0]
    for tag in ("[NAME]", "[EMAIL]", "[PHONE]", "[URL]"):
        assert tag in a.masked, a.masked
    assert "Django" in a.masked and "Kafka" in a.masked  # spaCy tags these as PERSON — never mask a skill
    assert "Aarav" not in a.masked and "98765" not in a.masked


def test_normalize_text() -> None:
    counter: Counter[str] = Counter()
    out = normalize_text("Py​thon ﬁnance “quoted” – Pythоn", counter)
    assert out == 'Python finance "quoted" - Python'
    assert counter["zero_width"] == 1 and counter["confusable_words"] == 1
    assert canonical("Design, build & maintain REST APIs!") == "design build maintain rest apis"


def test_sentence_split(nlp_engine) -> None:
    assert len(nlp_engine.sentences("Built APIs. Deployed to AWS. Wrote tests.")) == 3


def _units(fixtures_dir: Path, name: str, nlp_engine):
    settings = Settings()
    data = (fixtures_dir / name).read_bytes()
    doc = extract(data, name.rsplit(".", 1)[1], file_name=name, sha256="0" * 64, settings=settings)
    apply_visibility(doc, settings)
    return preprocess_document(doc, nlp_engine, candidate_name=extract_contact(doc).name)


def test_sections_and_units_from_pdf(fixtures_dir: Path, nlp_engine) -> None:
    units = _units(fixtures_dir, "clean.pdf", nlp_engine)
    sections = {u.section for u in units if not u.is_heading}
    assert {"summary", "skills", "experience", "projects", "education", "certifications"} <= sections
    by_text = {u.text: u for u in units}
    wrapped = (
        "Packaged the microservices as Docker images and ran them on AWS EC2 behind an Nginx load balancer."
    )
    assert (
        wrapped in by_text and by_text[wrapped].section == "experience"
    )  # two PDF lines merged into one bullet
    assert by_text["AARAV SHARMA"].excluded == "contact"
    assert any(u.excluded == "contact" and "@example.com" in u.text for u in units)
    assert by_text["Languages: Python, JavaScript, SQL"].is_skill_list
    assert sum(1 for u in units if u.is_heading) == 6


def test_docx_and_pdf_give_the_same_units(fixtures_dir: Path, nlp_engine) -> None:
    def scored(us):
        return {u.text.lower() for u in us if not u.is_heading and not u.excluded}

    pdf, docx = (
        scored(_units(fixtures_dir, "clean.pdf", nlp_engine)),
        scored(_units(fixtures_dir, "clean.docx", nlp_engine)),
    )
    assert len(pdf ^ docx) <= 2, pdf ^ docx


def test_sidebar_skill_labels_are_not_sections(fixtures_dir: Path, nlp_engine) -> None:
    units = _units(fixtures_dir, "dark_sidebar_legit.pdf", nlp_engine)
    assert not any(
        u.is_heading and u.section == "languages" and u.text == "Languages"
        for u in units
        if u.section == "skills"
    )
    assert "Python, JavaScript, SQL" in {u.text for u in units}


def test_personal_details_never_scored(nlp_engine) -> None:
    import io

    from docx import Document

    d = Document()
    for line in ["Riya Patil", "PERSONAL DETAILS", "Date of Birth: 01/02/2001", "Gender: Female", "Marital Status: Single",
                 "Father's Name: Suresh Patil", "SKILLS", "Python, Django, PostgreSQL, Docker"]:  # fmt: skip
        p = d.add_paragraph()
        r = p.add_run(line)
        r.bold = line.isupper()
    buf = io.BytesIO()
    d.save(buf)
    settings = Settings()
    doc = extract(buf.getvalue(), "docx", file_name="p.docx", sha256="0" * 64, settings=settings)
    apply_visibility(doc, settings)
    units = preprocess_document(doc, nlp_engine, candidate_name="Riya Patil")
    personal = [u for u in units if u.excluded == "personal_details"]
    assert {u.text.split(":")[0] for u in personal} >= {
        "Date of Birth",
        "Gender",
        "Marital Status",
        "Father's Name",
    }
