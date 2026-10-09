"""Integrity detectors: each fixture raises its expected flags, legit fixtures raise none, every flag has evidence."""

from __future__ import annotations

import io
from datetime import date
from pathlib import Path
from typing import Any

import numpy as np
import pytest

from app.domain.enums import ResumeStatus, Severity, SystemDecision
from app.domain.flags import Flag
from app.services.integrity.base import all_detectors
from app.services.integrity.risk import risk_score
from app.services.pipeline import AnalysisResult, analyze_document, prepare_job

pytestmark = pytest.mark.slow

EXPECTED: dict[str, set[str]] = {
    # hidden text
    "white_text_jd.pdf": {"LOW_CONTRAST", "JD_BULK_COPY", "VERBATIM_JD_LINE", "SCORE_INFLATION_ATTEMPT"},
    "same_color_box.pdf": {"LOW_CONTRAST"},
    "tiny_font.pdf": {"TINY_FONT"},
    "squeezed_text.pdf": {"TINY_FONT"},
    "low_opacity.pdf": {"LOW_OPACITY", "SCORE_INFLATION_ATTEMPT"},
    "render_mode_invisible.pdf": {"INVISIBLE_RENDER_MODE", "SCORE_INFLATION_ATTEMPT"},
    "off_page.pdf": {"OFF_PAGE", "SCORE_INFLATION_ATTEMPT"},
    "behind_image.pdf": {"COVERED"},
    "hidden_layer.pdf": {"HIDDEN_LAYER", "SCORE_INFLATION_ATTEMPT"},
    "stacked_text.pdf": {"STACKED_TEXT"},
    "clipped_text.pdf": {"NOT_RENDERED_VISIBLY"},
    "hidden_vanish.docx": {"DOCX_HIDDEN", "SCORE_INFLATION_ATTEMPT"},
    "white_text.docx": {"DOCX_HIDDEN"},
    "tiny_font.docx": {"DOCX_HIDDEN"},
    "comment_stuffing.docx": {"DOCX_HIDDEN"},
    "small_font_suspicious.pdf": {"TINY_FONT"},
    # content tricks
    "verbatim_jd_lines.pdf": {"VERBATIM_JD_LINE"},
    "near_verbatim_jd.pdf": {"NEAR_VERBATIM_JD"},
    "jd_bulk_copy.pdf": {"JD_BULK_COPY", "VERBATIM_JD_LINE", "SCORE_INFLATION_ATTEMPT"},
    "jd_voice_phrases.pdf": {"JD_VOICE_PHRASES"},
    "keyword_stuffing.pdf": {"KEYWORD_STUFFING", "KEYWORD_DUMP", "REPEATED_CONTENT"},
    "prompt_injection_hidden.pdf": {"PROMPT_INJECTION", "LOW_CONTRAST"},
    "injection_paraphrase.pdf": {"PROMPT_INJECTION_PARAPHRASE", "LOW_CONTRAST"},
    "zero_width.pdf": {"INVISIBLE_UNICODE"},
    "homoglyph.pdf": {"HOMOGLYPHS"},
    "metadata_stuffing.pdf": {"METADATA_STUFFING"},
    "timeline_bad.pdf": {"TIMELINE_INCONSISTENT", "TIMELINE_FUTURE_DATE"},
    # files
    "encrypted.pdf": {"ENCRYPTED_FILE"},
    "javascript.pdf": {"MALICIOUS_PDF_FEATURES"},
    "scanned_image_only.pdf": {"SCANNED_DOCUMENT"},
}
LEGIT = [
    "clean.pdf",
    "clean.docx",
    "dark_sidebar_legit.pdf",
    "tailored_legit.pdf",
    "image_header_legit.pdf",
    "unrelated.pdf",
]
TODAY = date(2026, 10, 9)


@pytest.fixture(scope="module")
def job(resources, sample_jd: str):
    return prepare_job("Backend Engineer (Python)", sample_jd, resources)


@pytest.fixture(scope="module")
def analysed(resources, job, fixtures_dir: Path):
    cache: dict[str, AnalysisResult] = {}

    def run(name: str) -> AnalysisResult:
        if name not in cache:
            cache[name] = analyze_document(
                (fixtures_dir / name).read_bytes(), name, job, resources, today=TODAY
            )
        return cache[name]

    return run


def codes(result: AnalysisResult) -> set[str]:
    return {f.code for f in result.flags}


@pytest.mark.parametrize(("name", "expected"), sorted(EXPECTED.items()))
def test_fixture_raises_expected_flags(analysed, name: str, expected: set[str]) -> None:
    got = codes(analysed(name))
    assert expected <= got, f"missing {expected - got}; got {got}"


@pytest.mark.parametrize("name", LEGIT)
def test_legit_fixture_raises_no_flags(analysed, name: str) -> None:
    result = analysed(name)
    assert result.flags == [], [(f.code, f.message) for f in result.flags]
    assert result.risk_score == 0 and result.risk_level is None


@pytest.mark.parametrize("name", sorted(EXPECTED))
def test_every_flag_carries_evidence(analysed, name: str) -> None:
    for flag in analysed(name).flags:
        ev = flag.evidence
        assert flag.message and flag.detector and ev.location, flag
        if ev.location in ("visible", "hidden", "mixed"):
            if name.endswith(".pdf"):
                assert ev.page is not None and ev.bbox is not None, flag
            else:  # DOCX has no fixed layout: the evidence names the paragraph / comment instead of a bbox
                assert ev.page is not None or "position" in ev.details, flag
            assert ev.snippet, flag
        assert ev.measured_value is not None, flag
        if flag.code not in ("ZIP_BOMB",):
            assert ev.threshold is not None, flag


@pytest.mark.parametrize(
    ("name", "code", "severity"),
    [
        ("white_text_jd.pdf", "LOW_CONTRAST", Severity.CRITICAL),  # hidden text containing JD sentences
        ("prompt_injection_hidden.pdf", "LOW_CONTRAST", Severity.HIGH),  # hidden, but no JD content
        ("prompt_injection_hidden.pdf", "PROMPT_INJECTION", Severity.CRITICAL),
        ("injection_paraphrase.pdf", "PROMPT_INJECTION_PARAPHRASE", Severity.HIGH),
        ("white_text_jd.pdf", "SCORE_INFLATION_ATTEMPT", Severity.HIGH),
        ("jd_bulk_copy.pdf", "JD_BULK_COPY", Severity.CRITICAL),
        ("verbatim_jd_lines.pdf", "VERBATIM_JD_LINE", Severity.HIGH),
        ("near_verbatim_jd.pdf", "NEAR_VERBATIM_JD", Severity.MEDIUM),
        ("small_font_suspicious.pdf", "TINY_FONT", Severity.MEDIUM),  # 5 pt: readable but suspicious
        ("tiny_font.pdf", "TINY_FONT", Severity.CRITICAL),
        ("javascript.pdf", "MALICIOUS_PDF_FEATURES", Severity.CRITICAL),
        ("timeline_bad.pdf", "TIMELINE_INCONSISTENT", Severity.LOW),
        ("homoglyph.pdf", "HOMOGLYPHS", Severity.HIGH),
    ],
)
def test_severity(analysed, name: str, code: str, severity: Severity) -> None:
    flag = next(f for f in analysed(name).flags if f.code == code)
    assert Severity(flag.severity) == severity


def test_hidden_text_evidence_details(analysed) -> None:
    flag = next(f for f in analysed("white_text_jd.pdf").flags if f.code == "LOW_CONTRAST")
    ev = flag.evidence
    assert ev.measured_value is not None and ev.measured_value < 1.1 and ev.threshold == 1.5
    assert ev.details["jd_sentences"] and ev.details["confidence"] == "high"
    assert "containing JD sentences" in flag.message
    assert len(ev.regions) >= 5


def test_verbatim_evidence_says_hidden_or_visible(analysed) -> None:
    hidden = next(f for f in analysed("white_text_jd.pdf").flags if f.code == "VERBATIM_JD_LINE")
    visible = next(f for f in analysed("verbatim_jd_lines.pdf").flags if f.code == "VERBATIM_JD_LINE")
    assert hidden.evidence.location == "hidden" and visible.evidence.location == "visible"


@pytest.mark.parametrize(
    ("name", "decision"),
    [
        ("clean.pdf", SystemDecision.SHORTLISTED),
        ("unrelated.pdf", SystemDecision.NOT_SHORTLISTED),
        ("white_text_jd.pdf", SystemDecision.FLAGGED),
        ("prompt_injection_hidden.pdf", SystemDecision.FLAGGED),
        ("javascript.pdf", SystemDecision.FLAGGED),
        ("encrypted.pdf", SystemDecision.NEEDS_REVIEW),
        ("scanned_image_only.pdf", SystemDecision.NEEDS_REVIEW),
        ("tailored_legit.pdf", SystemDecision.NEEDS_REVIEW),  # borderline clean score, no flags
        ("zero_width.pdf", SystemDecision.NEEDS_REVIEW),  # ≥ 70 but a high flag
    ],
)
def test_decision(analysed, name: str, decision: SystemDecision) -> None:
    assert analysed(name).decision.system == decision


def test_status_for_unreadable_files(analysed) -> None:
    assert analysed("encrypted.pdf").status == ResumeStatus.UNPARSEABLE
    assert analysed("scanned_image_only.pdf").status == ResumeStatus.UNPARSEABLE
    assert analysed("clean.pdf").status == ResumeStatus.DONE


def test_risk_score_is_capped_sum() -> None:
    from app.core.config import RiskSettings
    from app.domain.flags import Evidence

    def f(sev: Severity) -> Flag:
        return Flag("X", sev, "t", "m", Evidence())

    assert risk_score([f(Severity.LOW), f(Severity.MEDIUM)], RiskSettings()) == (13, Severity.MEDIUM)
    assert risk_score([f(Severity.CRITICAL)] * 3, RiskSettings()) == (100, Severity.CRITICAL)
    assert risk_score([], RiskSettings()) == (0, None)


def test_detectors_are_documented_and_registered() -> None:
    names = [d.name for d in all_detectors()]
    assert len(names) == len(set(names)) == 9
    for det in all_detectors():
        assert det.__doc__ and len(det.__doc__.split()) > 20, det.name


# ------------------------------------------------------------------------------------------ synthetic documents


def docx_bytes(paragraphs: list[str]) -> bytes:
    from docx import Document

    d = Document()
    d.add_paragraph("Kiran Rao").runs[0].bold = True
    d.add_paragraph("kiran.rao@example.com | +91 90000 11111")
    for text in paragraphs:
        d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def run_docx(paragraphs: list[str], job, resources, **kw: Any) -> AnalysisResult:
    return analyze_document(docx_bytes(paragraphs), "synthetic.docx", job, resources, today=TODAY, **kw)


def test_legit_lines_are_not_injection_or_voice(job, resources) -> None:
    result = run_docx(
        [
            "EXPERIENCE",
            "Built an AI assistant that ranks support tickets by urgency using machine learning.",
            "Developed a resume screening tool that scores candidates against job descriptions.",
            "Selected as the best intern of the year and recommended by the client for excellent delivery.",
            "Ranked in the top 5% of 20,000 participants in a national coding contest.",
        ],
        job,
        resources,
    )
    assert not codes(result) & {"PROMPT_INJECTION", "PROMPT_INJECTION_PARAPHRASE", "JD_VOICE_PHRASES"}


def test_labelled_tech_list_under_a_project_is_not_a_dump(job, resources) -> None:
    result = run_docx(
        [
            "PROJECTS",
            "Payments Gateway - built a payout service for small merchants.",
            "Tech stack: Python, FastAPI, PostgreSQL, Redis, Docker, AWS, Git, Kafka, Kubernetes, pytest, GitHub Actions, Grafana",
        ],
        job,
        resources,
    )
    assert "KEYWORD_DUMP" not in codes(result)


def test_oversized_skill_list(job, resources) -> None:
    items = ", ".join(f"Tool{i}" for i in range(60))
    result = run_docx(["SKILLS", f"Tools: {items}"], job, resources)
    flag = next(f for f in result.flags if f.code == "KEYWORD_DUMP")
    assert flag.evidence.measured_value == 60


def test_garbled_text(job, resources) -> None:
    garbled = " ".join(["xkcdvb", "qwrtpl", "zxcvbn", "mnbvcx", "", "lkjhgf", "rtzwqp"] * 12)
    result = run_docx(["SUMMARY", garbled, "Python developer with Django experience."], job, resources)
    flag = next(f for f in result.flags if f.code == "GARBLED_TEXT")
    assert Severity(flag.severity) == Severity.LOW and flag.evidence.measured_value >= 0.4


def test_expected_and_present_dates_are_fine(job, resources) -> None:
    result = run_docx(
        [
            "EDUCATION",
            "B.E. Computer Engineering, Pune University (2023 - 2027, expected)",
            "EXPERIENCE",
            "Intern, Acme (Jun 2025 - Present)",
        ],
        job,
        resources,
    )
    assert not codes(result) & {"TIMELINE_INCONSISTENT", "TIMELINE_FUTURE_DATE"}


def test_short_alt_text_is_not_flagged(job, resources, fixtures_dir: Path) -> None:
    """A profile-photo alt text ("Photo of Aarav") is normal; only long or keyword-heavy alt text is flagged."""
    import re
    import zipfile

    src = zipfile.ZipFile(fixtures_dir / "clean.docx")
    xml = src.read("word/document.xml").decode()
    drawing = (
        '<w:p><w:r><w:drawing><wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        '<wp:docPr id="7" name="Picture 7" descr="Photo of Aarav"/></wp:inline></w:drawing></w:r></w:p>'
    )
    xml = re.sub(r"<w:body>", "<w:body>" + drawing, xml, count=1)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            dst.writestr(
                item, xml.encode() if item.filename == "word/document.xml" else src.read(item.filename)
            )
    result = analyze_document(out.getvalue(), "alt.docx", job, resources, today=TODAY)
    assert result.flags == []
    assert any(s.attrs.get("docx_hidden") == "alt_text" for s in result.doc.spans)  # type: ignore[union-attr]


class FakeIndex:
    def __init__(self, sha=None, near=None, contact=None):
        self.sha, self.near, self.contact = sha or [], near or [], contact or []

    def same_sha256(self, sha256: str) -> list[dict[str, Any]]:
        return self.sha

    def nearest_resumes(self, embedding: np.ndarray, limit: int = 3) -> list[dict[str, Any]]:
        return self.near

    def same_contact(self, email: str | None, phone: str | None) -> list[dict[str, Any]]:
        return self.contact


def test_duplicate_detectors(job, resources, fixtures_dir: Path) -> None:
    data = (fixtures_dir / "clean.pdf").read_bytes()
    index = FakeIndex(
        sha=[{"resume_id": "r1", "file_name": "clean (1).pdf"}],
        near=[
            {
                "resume_id": "r2",
                "file_name": "copy.pdf",
                "candidate_name": "Someone Else",
                "similarity": 0.985,
            }
        ],
        contact=[{"resume_id": "r3", "candidate_name": "Rahul Mehta", "email": "aarav.sharma@example.com"}],
    )
    result = analyze_document(data, "clean.pdf", job, resources, duplicate_index=index, today=TODAY)
    got = {f.code: f for f in result.flags}
    assert {"DUPLICATE_FILE", "NEAR_DUPLICATE_RESUME", "IDENTITY_CONFLICT"} <= set(got)
    assert got["NEAR_DUPLICATE_RESUME"].evidence.threshold == 0.97
    assert Severity(got["NEAR_DUPLICATE_RESUME"].severity) == Severity.MEDIUM  # a different person
    resubmission = FakeIndex(
        near=[
            {"resume_id": "r5", "file_name": "old.pdf", "candidate_name": "Aarav Sharma", "similarity": 0.99}
        ]
    )
    result2 = analyze_document(data, "clean.pdf", job, resources, duplicate_index=resubmission, today=TODAY)
    near = next(f for f in result2.flags if f.code == "NEAR_DUPLICATE_RESUME")
    assert Severity(near.severity) == Severity.LOW and "same candidate" in near.message
    assert Severity(got["IDENTITY_CONFLICT"].severity) == Severity.HIGH
    same_person = FakeIndex(contact=[{"resume_id": "r4", "candidate_name": "Aarav Sharma"}])
    result = analyze_document(data, "clean.pdf", job, resources, duplicate_index=same_person, today=TODAY)
    assert "IDENTITY_CONFLICT" not in codes(result)  # same email + same name = the same person applying again


def test_failed_file_never_breaks(job, resources) -> None:
    result = analyze_document(b"%PDF-1.7\n this is not really a pdf", "broken.pdf", job, resources)
    assert result.status == ResumeStatus.FAILED and result.status_reason
